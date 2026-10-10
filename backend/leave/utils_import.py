import datetime as dt
from io import BytesIO
import pandas as pd
from django.db import transaction
from django.utils import timezone
from django.utils.dateparse import parse_date
from company.models import CompanyGroup
from leave.models import ImportedLeaveBalance, Leave, LeaveType
from user.models import CustomUser
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


STAFF_ALIASES = ('staff id', 'staffid', 'staff_code', 'emp_code',
                 'biometric id', 'biometric', 'biometric_id')
# 'Staff ID' is the authoritative column; 'Biometric ID' is only a fallback
# used when a sheet has no 'Staff ID' column at all.
STAFF_ID_ALIASES = ('staff id', 'staffid', 'staff_code', 'emp_code')
BIOMETRIC_ALIASES = ('biometric id', 'biometric', 'biometric_id')


def _find_staff_column(columns):
    """Locate the staff column, preferring 'Staff ID' over 'Biometric ID'."""
    col = _find_column(columns, STAFF_ID_ALIASES)
    if col:
        return col
    return _find_column(columns, BIOMETRIC_ALIASES)


# ---------------------------------------------------------------------------
# File reading helpers
# ---------------------------------------------------------------------------

def _read_raw(file_path):
    """Read the uploaded file into a raw DataFrame (header=None).
    Returns (df, error_message)."""
    name = (getattr(file_path, 'name', '') or '').lower()
    try:
        if name.endswith('.csv'):
            return pd.read_csv(file_path, header=None), None
        if name.endswith(('.xlsx', '.xlsm')):
            return pd.read_excel(file_path, header=None, engine='openpyxl'), None
        if name.endswith('.xls'):
            return None, ".xls is not supported. Save the file as .xlsx and retry."
        return None, "Unsupported file type. Upload a .xlsx or .csv file."
    except Exception as e:
        return None, f"Error reading file: {str(e)}"


def _to_float(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    try:
        return float(str(value).replace(',', '').strip())
    except (ValueError, TypeError):
        return None


# ---------------------------------------------------------------------------
# Bulk loaders (avoid N+1 - the remote DB costs ~340ms per query)
# ---------------------------------------------------------------------------

def _load_user_maps(company):
    """biometric_id -> user_id, preferring a parent_company match over M2M.
    2 queries total instead of 2 per file row."""
    def _map(qs):
        return {
            str(b): uid
            for uid, b in qs.exclude(biometric_id__isnull=True).exclude(biometric_id='')
            .values_list('id', 'biometric_id')
        }
    parent = _map(CustomUser.objects.filter(parent_company=company))
    m2m = _map(CustomUser.objects.filter(company=company))
    return parent, m2m


def _resolve_user_id(maps, staff_id):
    parent, m2m = maps
    return parent.get(staff_id) or m2m.get(staff_id)


def _load_existing_balances(company, year):
    """(user_id, leave_type_id) -> existing row. 1 query."""
    return {
        (row.user_id, row.leave_type_id): row
        for row in ImportedLeaveBalance.objects.filter(company=company, year=year)
        .only('id', 'user_id', 'leave_type_id', 'snapshot_date',
              'opening_balance', 'expiry_date')
    }


def _load_leave_types(company):
    """short_name -> LeaveType (all company leave types). 1 query."""
    return {
        lt.short_name: lt
        for lt in LeaveType.objects.filter(company=company)
    }


def _flush(pending, existing, company, year, snapshot_date, expiry_date, dry_run):
    """Write pending balances in bulk. pending: {(user_id, leave_type_id): value}."""
    created = updated = 0
    to_create, to_update = [], []

    for (user_id, leave_type_id), value in pending.items():
        row = existing.get((user_id, leave_type_id))
        if row is None:
            created += 1
            if not dry_run:
                to_create.append(ImportedLeaveBalance(
                    user_id=user_id,
                    leave_type_id=leave_type_id,
                    company=company,
                    year=year,
                    snapshot_date=snapshot_date,
                    opening_balance=value,
                    expiry_date=expiry_date,
                ))
        else:
            updated += 1
            if not dry_run:
                row.snapshot_date = snapshot_date
                row.opening_balance = value
                # Omit expiry when not provided so a re-import never wipes
                # an expiry that was already set on this row.
                if expiry_date is not None:
                    row.expiry_date = expiry_date
                row.imported_at = timezone.now()
                to_update.append(row)

    if not dry_run:
        if to_create:
            ImportedLeaveBalance.objects.bulk_create(to_create, batch_size=500)
        if to_update:
            fields = ['snapshot_date', 'opening_balance', 'imported_at']
            if expiry_date is not None:
                fields.append('expiry_date')
            ImportedLeaveBalance.objects.bulk_update(
                to_update,
                fields,
                batch_size=500,
            )
    return created, updated


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def import_balances_from_excel(file_path, company, year, snapshot_date_str,
                               expiry_date=None, dry_run=False):
    """
    Parses an Excel/CSV file containing opening leave balances and imports them
    into ImportedLeaveBalance (one row per user + leave_type + company + year).

    Expected format - a simple sheet with a header row of
    [Staff ID, Employee Name, <LeaveType.short_name>, ...]. Columns whose
    header matches one of the company's LeaveType short_names are imported.

    Matching is on CustomUser.biometric_id (the "Staff ID" column).

    expiry_date: optional date (or YYYY-MM-DD string). The import stays in use
    while today <= expiry_date. If omitted on a re-import, existing expiry dates
    are kept. New rows get expiry_date=None (never expires).

    dry_run=True -> parse, match and report what WOULD happen, without writing
    any LeaveType or ImportedLeaveBalance rows.
    """
    snapshot_date = parse_date(snapshot_date_str)
    if not snapshot_date:
        return {'success': False, 'message': "Invalid snapshot date format (expected YYYY-MM-DD)."}

    if expiry_date is not None and not isinstance(expiry_date, dt.date):
        expiry_date = parse_date(str(expiry_date))
    if expiry_date is not None and not isinstance(expiry_date, dt.date):
        return {'success': False, 'message': "Invalid expiry date format (expected YYYY-MM-DD)."}

    raw, err = _read_raw(file_path)
    if err:
        return {'success': False, 'message': err}
    if raw is None or raw.empty:
        return {'success': False, 'message': "The uploaded file is empty."}

    return _import_standard(raw, company, year, snapshot_date, expiry_date, dry_run)


# ---------------------------------------------------------------------------
# Simple sheet format
# ---------------------------------------------------------------------------

def _import_standard(raw, company, year, snapshot_date, expiry_date=None, dry_run=False):
    df = raw.copy()
    df.columns = [str(c).strip() for c in df.iloc[0].tolist()]
    df = df.iloc[1:].reset_index(drop=True)

    if df.empty:
        return {'success': False, 'message': "The uploaded file is empty."}

    staff_col = _find_staff_column(df.columns)
    if staff_col is None:
        return {'success': False,
                'message': "No 'Staff ID' column found. Expected a column named 'Staff ID'."}

    available_leave_types = _load_leave_types(company)
    if not available_leave_types:
        return {'success': False,
                'message': "No leave types exist for this company. "
                           "Create leave types for the company first."}

    # Columns are matched against the company's leave type short_name OR full
    # name (case-insensitive), e.g. 'CL' or 'Casual Leave'.
    type_by_header = {}
    for lt in available_leave_types.values():
        if lt.short_name:
            type_by_header.setdefault(_canon(lt.short_name), lt)
        if lt.leave_type:
            type_by_header.setdefault(_canon(lt.leave_type), lt)

    matched_columns = [c for c in df.columns if _canon(c) in type_by_header]
    if not matched_columns:
        return {'success': False,
                'message': "No Excel columns match this company's leave types. "
                           "Re-download the current template from the app. "
                           "Expected columns like "
                           f"{[lt.leave_type for lt in available_leave_types.values()][:8]}..."}

    results = {'updated': 0, 'created': 0, 'rows_processed': 0,
               'columns_found': matched_columns, 'unmatched': [], 'errors': [],
               'dry_run': dry_run,
               'expiry_date': expiry_date.isoformat() if expiry_date else None}

    user_maps = _load_user_maps(company)
    existing = _load_existing_balances(company, year)
    pending = {}

    with transaction.atomic():
        for index, row in df.iterrows():
            staff_id_val = row.get(staff_col)
            if staff_id_val is None or pd.isna(staff_id_val):
                continue
            staff_id = str(staff_id_val).strip()
            if staff_id.endswith('.0'):
                staff_id = staff_id[:-2]
            if not staff_id or staff_id.lower() == 'nan':
                continue

            user_id = _resolve_user_id(user_maps, staff_id)
            if not user_id:
                results['unmatched'].append(staff_id)
                continue

            results['rows_processed'] += 1

            for col_name in matched_columns:
                balance_val = _to_float(row.get(col_name))
                if balance_val is None:
                    continue
                pending[(user_id, type_by_header[_canon(col_name)].id)] = balance_val

        results['created'], results['updated'] = _flush(
            pending, existing, company, year, snapshot_date, expiry_date, dry_run)

    if results['unmatched']:
        results['errors'].append(
            f"{len(results['unmatched'])} staff IDs not found: "
            f"{results['unmatched'][:15]}{'...' if len(results['unmatched']) > 15 else ''}"
        )
    return {'success': True, 'results': results}


# ---------------------------------------------------------------------------
# Leave-record sheet (one row = one past leave record, saved into the Leave table)
# ---------------------------------------------------------------------------

LEAVE_TYPE_ALIASES = ('leave type', 'leave type name', 'leave name', 'leavetype',
                      'leave', 'type of leave')
FROM_DATE_ALIASES = ('from date', 'start date', 'leave from', 'leave from date')
TO_DATE_ALIASES = ('to date', 'end date', 'leave to', 'leave to date', 'till date')
DAYS_ALIASES = ('days taken', 'days', 'no of days', 'no. of days', 'number of days',
                'total days')
CHOICE_ALIASES = ('leave choice', 'duration', 'half day', 'day type', 'is half day', 'half')
STATUS_ALIASES = ('status', 'approval status', 'leave status')
REASON_ALIASES = ('reason', 'remarks', 'custom reason', 'note', 'notes', 'description')


def _canon(value):
    """Lower-case, collapse underscores to spaces, drop '*' flags and trim."""
    return ' '.join(str(value).strip().lower()
                    .replace('_', ' ').replace('*', '').split())


def _find_column(columns, aliases):
    alias_set = set(_canon(a) for a in aliases)
    for col in columns:
        if _canon(col) in alias_set:
            return col
    return None


def _parse_date_cell(value):
    """Parse an Excel date cell (datetime, date, serial number or string)."""
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (int, float)):
        # Excel serial date (days since 1899-12-30); ignore absurd values.
        serial = int(value)
        if 1 <= serial <= 60000:
            return dt.date(1899, 12, 30) + dt.timedelta(days=serial)
        return None
    s = str(value).strip()
    for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%Y/%m/%d', '%d/%m/%Y', '%d.%m.%Y'):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _parse_choice(value):
    s = _canon(value)
    if s in ('h', 'half', 'half day', 'half-day', '0.5', '0.5 day'):
        return 'H'
    return 'F'


def _parse_status(value):
    s = _canon(value)
    if s:
        first = s[0]
        if first == 'p':
            return 'P'
        if first == 'r':
            return 'R'
        if first == 'c':
            return 'C'
    return 'A'


def import_leave_records_from_excel(file_path, company, dry_run=False):
    """
    Parses a sheet where one row = one past leave record and inserts the matched
    rows directly into the Leave table (no ImportedLeaveBalance involved).

    Required columns (headers matched case-insensitively):
      Staff ID | Leave Type | From Date | To Date
    Optional columns:
      Days Taken, Leave Choice (half day / full day), Status (A/P/R/C), Reason.

    - Staff ID is matched against CustomUser.biometric_id.
    - Leave Type is matched against the company's LeaveType short_name or name.
    - days_taken = (To - From) + 1 (halved for a half day) unless Days Taken is
      given in the sheet.
    - Status defaults to 'A' (approved).

    Rows with missing/invalid data are reported in results['errors'] and skipped;
    nothing is written when dry_run=True.
    """
    raw, err = _read_raw(file_path)
    if err:
        return {'success': False, 'message': err}
    if raw is None or raw.empty:
        return {'success': False, 'message': "The uploaded file is empty."}

    df = raw.copy()
    df.columns = [str(c).strip() for c in df.iloc[0].tolist()]
    df = df.iloc[1:].reset_index(drop=True)
    if df.empty:
        return {'success': False, 'message': "The uploaded file is empty."}

    staff_col = _find_staff_column(df.columns)
    type_col = _find_column(df.columns, LEAVE_TYPE_ALIASES)
    from_col = _find_column(df.columns, FROM_DATE_ALIASES)
    to_col = _find_column(df.columns, TO_DATE_ALIASES)
    days_col = _find_column(df.columns, DAYS_ALIASES)
    choice_col = _find_column(df.columns, CHOICE_ALIASES)
    status_col = _find_column(df.columns, STATUS_ALIASES)
    reason_col = _find_column(df.columns, REASON_ALIASES)

    labels = {
        'Staff ID': staff_col, 'Leave Type': type_col,
        'From Date': from_col, 'To Date': to_col,
    }
    missing = [name for name, col in labels.items() if not col]
    if missing:
        return {'success': False,
                'message': "Missing required column(s): " + ", ".join(missing) +
                           ". The sheet must match the sample template."}

    results = {
        'created': 0, 'updated': 0, 'rows_processed': 0,
        'columns_found': [c for c in labels.values()] +
                         [c for c in [days_col, choice_col, status_col, reason_col] if c],
        'unmatched': [], 'created_leave_types': [], 'errors': [],
        'dry_run': dry_run, 'expiry_date': None,
    }

    user_maps = _load_user_maps(company)
    available_leave_types = _load_leave_types(company)
    by_short = {_canon(lt.short_name): lt for lt in available_leave_types.values() if lt.short_name}
    by_name = {_canon(lt.leave_type): lt for lt in available_leave_types.values() if lt.leave_type}

    def _resolve_leave_type(value):
        s = _canon(value)
        if not s:
            return None
        if s in by_short:
            return by_short[s]
        if s in by_name:
            return by_name[s]
        # Tolerate "casual leave" when only the short name "CL" exists.
        if s.endswith(' leave') and s[:-6] in by_short:
            return by_short[s[:-6]]
        return None

    to_create = []
    seen = set()

    with transaction.atomic():
        for index, row in df.iterrows():
            sid_val = row.get(staff_col)
            if sid_val is None or pd.isna(sid_val):
                continue
            sid = str(sid_val).strip()
            if sid.endswith('.0'):
                sid = sid[:-2]
            if not sid or sid.lower() == 'nan':
                continue

            user_id = _resolve_user_id(user_maps, sid)
            if not user_id:
                results['unmatched'].append(sid)
                continue

            lt = _resolve_leave_type(row.get(type_col))
            if not lt:
                results['errors'].append(
                    f"Row for {sid}: unknown leave type '{row.get(type_col)}' - skipped.")
                continue

            from_date = _parse_date_cell(row.get(from_col))
            if not from_date:
                results['errors'].append(
                    f"Row for {sid} ({lt.short_name or lt.leave_type}): "
                    f"invalid 'From Date' '{row.get(from_col)}' - skipped.")
                continue
            to_date = _parse_date_cell(row.get(to_col)) or from_date
            if to_date < from_date:
                results['errors'].append(
                    f"Row for {sid} ({lt.short_name or lt.leave_type}): "
                    f"'To Date' is before 'From Date' - skipped.")
                continue

            choice = _parse_choice(row.get(choice_col)) if choice_col else 'F'
            status = _parse_status(row.get(status_col)) if status_col else 'A'

            reason = None
            if reason_col:
                rv = row.get(reason_col)
                if rv is not None and not pd.isna(rv):
                    reason = str(rv).strip() or None

            days_taken = None
            if days_col:
                days_taken = _to_float(row.get(days_col))
            if days_taken is None:
                base = (to_date - from_date).days + 1
                days_taken = base * 0.5 if choice == 'H' else float(base)

            key = (user_id, lt.id, from_date, to_date, choice)
            if key in seen:
                continue
            seen.add(key)
            results['rows_processed'] += 1

            to_create.append(Leave(
                user_id=user_id,
                leave_type=lt,
                company=company,
                from_date=from_date,
                to_date=to_date,
                leave_choice=choice,
                status=status,
                days_taken=days_taken,
                custom_reason=reason,
            ))

        results['created'] = len(to_create)
        if to_create and not dry_run:
            Leave.objects.bulk_create(to_create, batch_size=500)

    if results['unmatched']:
        results['errors'].append(
            f"{len(results['unmatched'])} staff IDs not found: "
            f"{results['unmatched'][:15]}{'...' if len(results['unmatched']) > 15 else ''}"
        )
    return {'success': True, 'results': results}


# ---------------------------------------------------------------------------
# Sample template builders (openpyxl) - one template per company so the leave
# type dropdowns/lists match the company's actual LeaveType records.
# ---------------------------------------------------------------------------

_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_EXAMPLE_FONT = Font(italic=True, color="9E9E9E")
_THIN = Side(style="thin", color="B0BEC5")
_CELL_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_DATE_FMT = "DD-MM-YYYY"
_IDENTITY_HEADERS = [
    ("User ID *", 10), ("Biometric ID *", 14), ("Staff ID *", 14),
    ("Staff Name *", 24), ("Department *", 14),
]
_DAY_TYPES = ("Full day", "Half day")
_STATUSES = ("Approved", "Rejected", "Cancelled", "Pending")


def _company_leave_types(company):
    return list(LeaveType.objects.filter(company=company).order_by('id'))


def _style_header(ws):
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _CELL_BORDER
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30


def _add_dropdown(ws, formula, cell_range):
    dv = DataValidation(type="list", formula1=formula,
                        allow_blank=True, showErrorMessage=True)
    ws.add_data_validation(dv)
    dv.add(cell_range)


def _add_lists_sheet(wb, leave_types, departments=None):
    """Lists sheet used by the dropdowns (Leave Type / Day Type / Status / Department)."""
    ws = wb.create_sheet("Lists")
    ws["A1"] = "Leave Type"
    ws["B1"] = "Day Type"
    ws["C1"] = "Status"
    ws["D1"] = "Department"
    for cell in ws[1]:
        cell.font = Font(bold=True)
    row = 2
    for lt in leave_types:
        if lt.leave_type:
            ws.cell(row=row, column=1, value=lt.leave_type)
            row += 1
    for i, day in enumerate(_DAY_TYPES):
        ws.cell(row=2 + i, column=2, value=day)
    for i, st in enumerate(_STATUSES):
        ws.cell(row=2 + i, column=3, value=st)
    dept_row = 2
    for d in (departments or []):
        ws.cell(row=dept_row, column=4, value=d)
        dept_row += 1
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 22
    return {
        'leave_type': 1 + sum(1 for lt in leave_types if lt.leave_type),
        # Last row index of the Department list (start = row 2); 1 when empty.
        'dept': max(dept_row - 1, 1),
    }


def _company_departments(company):
    """Department names for the company (CompanyGroup)."""
    groups = list(CompanyGroup.objects.filter(company=company).order_by('group'))
    return [g.group or g.short_name or f"Department {g.id}" for g in groups]


def _company_staff(company):
    """Ordered, de-duplicated biometric ids of the company's staff."""
    ids = set()
    for qs in (CustomUser.objects.filter(parent_company=company),
               CustomUser.objects.filter(company=company)):
        for (bio,) in qs.exclude(biometric_id__isnull=True).exclude(biometric_id='') \
                        .values_list('biometric_id'):
            if bio is not None:
                ids.add(str(bio).strip())
    return sorted(ids)


def _add_staff_sheet(wb, staff_ids):
    """Hidden 'Staff' sheet so the Staff ID column can be a dropdown."""
    if not staff_ids:
        return None
    ws = wb.create_sheet("Staff")
    ws["A1"] = "Staff ID"
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, sid in enumerate(staff_ids, start=2):
        ws.cell(row=i, column=1, value=sid)
    ws.column_dimensions["A"].width = 16
    ws.sheet_state = "hidden"
    return 1 + len(staff_ids)


def _add_instructions(wb, lines):
    ins = wb.create_sheet("Instructions")
    for i, (text, bold) in enumerate(lines, start=1):
        cell = ins.cell(row=i, column=1, value=text)
        if bold:
            cell.font = Font(bold=True, size=12)
    ins.column_dimensions["A"].width = 115
    return ins


def build_dates_template(company):
    """
    One row per leave record. Header layout:
    User ID | Biometric ID | Staff ID | Staff Name | Department | Leave Type |
    From Date | To Date | Day Type | Days Taken | Reason | Status
    Parsed by import_leave_records_from_excel.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Leave History"

    headers = [h for h, _ in _IDENTITY_HEADERS] + [
        "Leave Type *", "From Date *", "To Date *", "Day Type *",
        "Days Taken", "Reason", "Status",
    ]
    widths = [w for _, w in _IDENTITY_HEADERS] + [16, 14, 14, 12, 12, 28, 12]
    ws.append(headers)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    _style_header(ws)

    # Clearly-marked example row (row 2) - users replace or delete it.
    ws.append([
        102, 5515, "SNG-0000", "EXAMPLE - replace this row", "Civil",
        "Casual Leave", dt.datetime(2026, 3, 10), dt.datetime(2026, 3, 12),
        "Full day", 3, "Delete this example row before uploading.", "Approved",
    ])
    for cell in ws[2]:
        if cell.value is not None:
            cell.font = _EXAMPLE_FONT
        cell.border = _CELL_BORDER
    ws["G2"].number_format = _DATE_FMT
    ws["H2"].number_format = _DATE_FMT

    # Date cells (G, H) carry a dd-mm-yyyy format for the client.
    for r in range(3, 1001):
        ws.cell(row=r, column=7).number_format = _DATE_FMT
        ws.cell(row=r, column=8).number_format = _DATE_FMT

    leave_types = _company_leave_types(company)
    departments = _company_departments(company)
    staff_ids = _company_staff(company)
    lists = _add_lists_sheet(wb, leave_types, departments)
    staff_last = _add_staff_sheet(wb, staff_ids)

    _add_dropdown(ws, f"Lists!$A$2:$A${lists['leave_type']}", "F2:F1000")
    _add_dropdown(ws, "Lists!$B$2:$B$3", "I2:I1000")
    _add_dropdown(ws, "Lists!$C$2:$C$5", "L2:L1000")
    if departments:
        _add_dropdown(ws, f"Lists!$D$2:$D${lists['dept']}", "E2:E1000")
    if staff_last:
        _add_dropdown(ws, f"Staff!$A$2:$A${staff_last}", "C2:C1000")

    _add_instructions(wb, [
        ("Staff Leave History - How to fill", True),
        ("", False),
        ("1. One row per leave application. Replace or delete the example row (row 2) "
         "before uploading.", False),
        ("2. Staff ID: pick the employee from the dropdown (hidden Staff list). "
         "Department can also be picked from the dropdown.", False),
        ("3. Leave Type: choose from the dropdown.", False),
        ("4. From/To dates: dd-mm-yyyy. 'To Date' is the same as 'From Date' for a one-day leave.",
         False),
        ("5. Day Type: 'Full day' or 'Half day' (a half day must be a single date).", False),
        ("6. Days Taken: leave blank to count every calendar day in the range "
         "(half day = 0.5).", False),
        ("7. Status: leave blank if the leave was approved. Reason is optional.", False),
        ("8. Uploaded rows are saved directly into the leave records - no snapshot or "
         "expiry date is involved.", False),
    ])
    return wb


def build_balance_template(company):
    """
    One row per employee, one column per company leave type (balance counts).
    Parsed by import_balances_from_excel (columns matched by type short name
    or full name).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Leave History"

    leave_types = [lt for lt in _company_leave_types(company) if lt.leave_type]
    headers = [h for h, _ in _IDENTITY_HEADERS] + [lt.leave_type for lt in leave_types]
    widths = [w for _, w in _IDENTITY_HEADERS] + [15] * len(leave_types)
    ws.append(headers)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    _style_header(ws)

    example = [102, 5515, "SNG-0000", "EXAMPLE - replace this row", "Civil"]
    for i in range(len(leave_types)):
        example.append(12 if i == 0 else (4.5 if i == 1 else None))
    ws.append(example)
    for cell in ws[2]:
        if cell.value is not None:
            cell.font = _EXAMPLE_FONT
        cell.border = _CELL_BORDER

    departments = _company_departments(company)
    staff_ids = _company_staff(company)
    lists = _add_lists_sheet(wb, leave_types, departments)
    staff_last = _add_staff_sheet(wb, staff_ids)
    if departments:
        _add_dropdown(ws, f"Lists!$D$2:$D${lists['dept']}", "E2:E1000")
    if staff_last:
        _add_dropdown(ws, f"Staff!$A$2:$A${staff_last}", "C2:C1000")

    _add_instructions(wb, [
        ("Staff Leave Balance Count - How to fill", True),
        ("", False),
        ("1. One row per employee. Replace or delete the example row (row 2) before "
         "uploading.", False),
        ("2. Staff ID: pick the employee from the dropdown (hidden Staff list). "
         "Department can also be picked from the dropdown.", False),
        ("3. Each column after 'Department' is one leave type. Enter the opening day "
         "balance for that employee in the matching column.", False),
        ("4. Leave a cell blank when the employee has no balance for that leave type.", False),
        ("5. The Year, Uploading Date and Expiry Date are entered in the app before "
         "uploading, not in this sheet.", False),
    ])
    return wb